from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "RoietKPS_User_Manual_TH.docx"
OUTPUT = ROOT / "output" / "RoietKPS_User_Manual_TH_Complete.docx"
ASSET_DIR = ROOT / "tmp" / "manual-tracking-assets"
CONTRACT_SCREENSHOT = ROOT / "output" / "playwright" / "tracking-contract-mobile.png"
DEBT_SCREENSHOT = ROOT / "output" / "playwright" / "tracking-debt-desktop.png"

GREEN = "185F49"
DARK = "111111"
MUTED = "5A6B63"
PALE = "F2F7F4"
WHITE = "FFFFFF"
LIGHT_BORDER = "D9D9D9"
FONT = "Sarabun"


def set_run_font(run, size=None, bold=None, color=DARK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, size=12, bold=False, color=DARK, space_after=5, line_spacing=1.08):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LIGHT_BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn("w:" + edge)
        element = borders.find(tag)
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + margin_name))
        if node is None:
            node = OxmlElement("w:" + margin_name)
            margins.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def prepare_crops():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with Image.open(CONTRACT_SCREENSHOT) as image:
        width, height = image.size
        image.crop((0, 300, width, min(960, height))).save(ASSET_DIR / "contract-overview.png")
        image.crop((0, 2000, width, min(2750, height))).save(ASSET_DIR / "contract-schedule.png")
    with Image.open(DEBT_SCREENSHOT) as image:
        width, height = image.size
        image.crop((190, 150, width, min(1200, height))).save(ASSET_DIR / "repayment-overview.png")
        image.crop((190, 920, width, min(2110, height))).save(ASSET_DIR / "follow-up-history.png")


def add_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_section_heading(document, number, title, description):
    number_p = document.add_paragraph()
    number_p.paragraph_format.space_after = Pt(0)
    number_p.paragraph_format.keep_with_next = True
    number_run = number_p.add_run(f"{number:02d}")
    set_run_font(number_run, size=12, bold=True, color=GREEN)

    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.keep_with_next = True
    heading_run = heading.add_run(title)
    set_run_font(heading_run, size=28, bold=True, color=DARK)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.paragraph_format.keep_with_next = True
    subtitle_run = subtitle.add_run(description)
    set_run_font(subtitle_run, size=13, bold=False, color=MUTED)


def add_steps(document, steps):
    for index, step in enumerate(steps, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.05)
        paragraph.paragraph_format.first_line_indent = Inches(-0.05)
        paragraph.paragraph_format.space_after = Pt(2)
        number_run = paragraph.add_run(f"{index}. ")
        set_run_font(number_run, size=13, bold=True, color=GREEN)
        text_run = paragraph.add_run(step)
        set_run_font(text_run, size=13, bold=False, color=DARK)


def add_caption(document, text):
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(7)
    run = caption.add_run(text)
    set_run_font(run, size=10, bold=False, color=MUTED)


def add_picture(document, path, width, title, description):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    shape = paragraph.add_run().add_picture(str(path), width=width)
    add_alt_text(shape, title, description)
    return shape


def replace_text(document, old, new):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == old:
            paragraph.text = new
            return paragraph
    return None


def replace_text_containing(document, needle, new):
    for paragraph in document.paragraphs:
        if needle in paragraph.text:
            paragraph.text = new
            return paragraph
    return None


def remove_paragraph_border(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def remove_style_border(style):
    properties = style.element.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def normalize_page_breaks(document):
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        breaks = [
            node
            for node in paragraph._p.xpath(".//w:br")
            if node.get(qn("w:type")) == "page"
        ]
        if not breaks:
            continue

        next_paragraph = next(
            (candidate for candidate in paragraphs[index + 1 :] if candidate.text.strip()),
            None,
        )
        if next_paragraph is not None:
            next_paragraph.paragraph_format.page_break_before = True

        if not paragraph.text.strip():
            parent = paragraph._element.getparent()
            if parent is not None:
                parent.remove(paragraph._element)


def resize_first_picture_after_heading(document, heading_text, width_inches):
    paragraphs = list(document.paragraphs)
    heading_index = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == heading_text),
        None,
    )
    if heading_index is None:
        return

    for paragraph in paragraphs[heading_index + 1 :]:
        inline_nodes = paragraph._p.xpath(".//wp:inline")
        if not inline_nodes:
            continue

        inline = inline_nodes[0]
        extent_nodes = inline.xpath("./wp:extent")
        if not extent_nodes:
            return

        extent = extent_nodes[0]
        original_width = int(extent.get("cx"))
        original_height = int(extent.get("cy"))
        new_width = int(Inches(width_inches))
        new_height = int(original_height * new_width / original_width)
        extent.set("cx", str(new_width))
        extent.set("cy", str(new_height))

        for picture_extent in inline.xpath(".//a:xfrm/a:ext"):
            picture_extent.set("cx", str(new_width))
            picture_extent.set("cy", str(new_height))
        return


def update_existing_content(document):
    document.core_properties.title = "คู่มือการใช้งานระบบสารสนเทศกองทุนพัฒนาสหกรณ์ RoietKPS"
    document.core_properties.subject = "การใช้งานระบบตั้งแต่สมัครสมาชิก คำขอกู้ สัญญา และการติดตามหลังอนุมัติ"
    document.core_properties.author = "สำนักงานสหกรณ์จังหวัดร้อยเอ็ด"
    document.core_properties.comments = "คู่มือสำหรับผู้ใช้งานระบบ RoietKPS"

    title = replace_text(document, "ระบบสารสนเทศกองทุนพัฒนาสหกรณ์", "ระบบสารสนเทศกองทุนพัฒนาสหกรณ์")
    if title is not None:
        title.style = document.styles["Title"]
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(2)
        remove_paragraph_border(title)
        format_paragraph(title, size=34, bold=True, color=DARK, space_after=2, line_spacing=1.0)

    cover_version = replace_text(document, "ฉบับเริ่มต้น  •  ตุลาคม 2569", "ฉบับใช้งานระบบ  |  ตุลาคม 2569")
    if cover_version is not None:
        format_paragraph(cover_version, size=14, bold=True, color=GREEN, space_after=4)

    cover_scope = replace_text_containing(
        document,
        "จัดทำเพื่อสนับสนุน",
        "ครอบคลุมการสมัครสมาชิก คำขอกู้ การตรวจเอกสาร การอนุมัติ การจัดทำสัญญา และการติดตามการชำระ",
    )
    if cover_scope is not None:
        format_paragraph(cover_scope, size=12, bold=False, color=MUTED, space_after=4)

    overview = replace_text(
        document,
        "คู่มือนี้อธิบายเส้นทางงานหลัก ตั้งแต่สมัครสมาชิกจนถึงการอนุมัติคำขอกู้และติดตามผล",
        "คู่มือนี้อธิบายเส้นทางงานหลัก ตั้งแต่สมัครสมาชิกจนถึงการอนุมัติ จัดทำสัญญา รับชำระ และปิดสัญญา",
    )
    if overview is not None:
        format_paragraph(overview, size=13, bold=False, color=MUTED, space_after=5)

    workflow = replace_text(
        document,
        "สมัครสมาชิก  →  Admin อนุมัติ  →  จัดทำคำขอกู้  →  ส่งเอกสาร  →  เจ้าหน้าที่ส่งเสริมตรวจ  →  เจ้าหน้าที่จังหวัดอนุมัติ",
        "สมัครสมาชิก  →  Admin อนุมัติ  →  จัดทำคำขอกู้และเอกสาร  →  ตรวจสองระดับ  →  อนุมัติ  →  จัดทำสัญญา  →  เบิกจ่าย  →  ติดตามชำระ  →  ปิดสัญญา",
    )
    if workflow is not None:
        workflow.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(workflow, size=13, bold=True, color=GREEN, space_after=7)

    for index, paragraph in enumerate(document.paragraphs[:-1]):
        text = paragraph.text.strip()
        if len(text) == 2 and text.isdigit():
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
            format_paragraph(paragraph, size=12, bold=True, color=GREEN, space_after=0)
            heading = document.paragraphs[index + 1]
            heading.style = document.styles["Heading 1"]
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.left_indent = Inches(0)
            heading.paragraph_format.right_indent = Inches(0)
            heading.paragraph_format.first_line_indent = Inches(0)
            format_paragraph(heading, size=28, bold=True, color=DARK, space_after=2, line_spacing=1.0)

    role_table = document.tables[0]
    role_updates = {
        "Admin": "อนุมัติผู้ใช้ กำหนดสิทธิ จัดการข้อมูลหลังอนุมัติ กติกาแจ้งเตือน และข้อมูลส่วนกลาง",
        "เจ้าหน้าที่จังหวัด": "ตรวจและอนุมัติคำขอ จัดทำสัญญา บันทึกการเบิกจ่าย ตารางชำระ และรับชำระ",
        "เจ้าหน้าที่ส่งเสริม": "ตรวจเอกสาร ให้ความเห็น ส่งกลับแก้ไข และบันทึกผลติดตามหน่วยงานที่รับผิดชอบ",
        "สหกรณ์ / กลุ่มเกษตรกร": "จัดทำข้อมูลและคำขอกู้ ส่งเอกสาร ติดตามสัญญา กำหนดชำระ และประวัติรับชำระ",
    }
    for row_index, row in enumerate(role_table.rows):
        if row_index == 0:
            continue
        role_name = row.cells[0].text.strip()
        if role_name in role_updates:
            row.cells[1].text = role_updates[role_name]

    status_table = document.tables[1]
    existing_statuses = {row.cells[0].text.strip() for row in status_table.rows[1:]}
    for status, meaning, next_actor in (
        ("รอจัดทำสัญญา", "คำขออนุมัติแล้วและรอบันทึกข้อมูลสัญญา", "เจ้าหน้าที่จังหวัด"),
        ("อยู่ระหว่างชำระ", "เบิกจ่ายแล้วและมีการติดตามตามตารางชำระ", "จังหวัดและเจ้าหน้าที่ส่งเสริม"),
        ("ค้างชำระ", "มียอดที่ครบกำหนดแต่ยังชำระไม่ครบ", "จังหวัดและเจ้าหน้าที่ส่งเสริม"),
        ("ชำระครบ", "เงินต้นและยอดตามตารางชำระครบแล้ว", "เจ้าหน้าที่จังหวัดหรือ Admin ปิดสัญญา"),
    ):
        if status in existing_statuses:
            continue
        cells = status_table.add_row().cells
        for cell, value in zip(cells, (status, meaning, next_actor)):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph, size=10.5, bold=False, color=DARK, space_after=0, line_spacing=1.0)
        if len(status_table.rows) % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, PALE)

    resize_first_picture_after_heading(document, "จัดทำและติดตามคำขอกู้", 5.70)
    resize_first_picture_after_heading(document, "ส่งเอกสารและแก้ไขไฟล์ที่ถูกส่งกลับ", 5.70)
    resize_first_picture_after_heading(document, "ติดตามงานผ่านการแจ้งเตือน", 5.70)


def append_tracking_sections(document):
    document.add_page_break()
    add_section_heading(
        document,
        13,
        "จัดทำสัญญาหลังอนุมัติ",
        "คำขอที่จังหวัดอนุมัติจะสร้างรายการรอจัดทำสัญญาโดยอัตโนมัติ",
    )

    picture_line = document.add_paragraph()
    picture_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = picture_line.add_run().add_picture(str(ASSET_DIR / "contract-overview.png"), width=Inches(2.15))
    picture_line.add_run("   ")
    second = picture_line.add_run().add_picture(str(ASSET_DIR / "contract-schedule.png"), width=Inches(2.15))
    add_alt_text(first, "ข้อมูลสัญญา", "หน้าข้อมูลสัญญา แสดงสถานะ เงินต้น และวันที่สำคัญ")
    add_alt_text(second, "ตารางชำระ", "ส่วนเพิ่มตารางชำระหลังบันทึกการเบิกจ่าย")
    add_caption(document, "หน้าสัญญาเงินกู้และส่วนเพิ่มตารางชำระบนหน้าจอมือถือ")

    add_steps(document, [
        "เปิดเมนูสัญญาเงินกู้ แล้วเลือกรายการที่มีสถานะรอจัดทำสัญญา",
        "บันทึกเลขที่สัญญา วันที่เริ่มสัญญา วันที่ครบกำหนด และเลื่อนสถานะตามลำดับงาน",
        "เมื่อลงนามครบถ้วน ให้บันทึกวันและเวลาเบิกจ่าย ระบบจะเปลี่ยนสัญญาเป็นอยู่ระหว่างชำระ",
        "ตรวจยอดเงินต้น อัตราดอกเบี้ย ระยะเวลา และเจ้าหน้าที่ส่งเสริมก่อนเพิ่มตารางชำระ",
    ])

    document.add_page_break()
    add_section_heading(
        document,
        14,
        "สร้างตารางชำระและบันทึกรับชำระ",
        "เจ้าหน้าที่จังหวัดบันทึกงวดตามเอกสารแนบท้ายสัญญาและหลักฐานการรับเงิน",
    )
    add_picture(
        document,
        ASSET_DIR / "repayment-overview.png",
        Inches(6.35),
        "ตารางชำระและแบบฟอร์มรับชำระ",
        "หน้าติดตามหนี้ แสดงงวดตามกำหนด ยอดรับชำระ และแบบฟอร์มบันทึกรับชำระ",
    )
    add_caption(document, "ตารางชำระ สถานะรายงวด และแบบฟอร์มบันทึกรับชำระ")
    add_steps(document, [
        "เปิดสัญญาที่อยู่ระหว่างชำระ แล้วเพิ่มเลขงวด วันครบกำหนด เงินต้น ดอกเบี้ย และค่าปรับ",
        "เมื่อได้รับชำระ ให้เปิดเมนูติดตามหนี้และเลือกสัญญา",
        "เลือกงวดที่เกี่ยวข้อง ระบบจะเติมยอดคงค้างของงวดให้อัตโนมัติ",
        "กรอกเลขที่ใบรับ วันเวลา ช่องทางรับชำระ และตรวจยอดก่อนบันทึก",
        "กรณีชำระบางส่วน สถานะงวดจะยังคงเป็นชำระบางส่วนจนกว่ายอดตามงวดจะครบ",
    ])

    document.add_page_break()
    add_section_heading(
        document,
        15,
        "บันทึกผลการติดตาม",
        "เจ้าหน้าที่บันทึกการโทร อีเมล หนังสือ หรือลงพื้นที่ พร้อมกำหนดวันติดตามครั้งถัดไป",
    )
    add_picture(
        document,
        ASSET_DIR / "follow-up-history.png",
        Inches(6.20),
        "ผลการติดตามและประวัติ",
        "แบบฟอร์มบันทึกผลติดตาม ประวัติรับชำระ และประวัติการติดต่อหน่วยงาน",
    )
    add_caption(document, "แบบฟอร์มผลการติดตามและประวัติการดำเนินงานของสัญญา")
    add_steps(document, [
        "เลือกงวดที่เกี่ยวข้อง หรือเลือกภาพรวมสัญญาเมื่อเป็นการติดตามทั่วไป",
        "ระบุวันเวลา ช่องทาง ผลการติดตาม และรายละเอียดที่จำเป็นต่อการดำเนินงานครั้งถัดไป",
        "หากต้องติดตามต่อ ให้กำหนดวันที่ครั้งถัดไป ระบบจะสร้างงานบนแดชบอร์ดของผู้บันทึก",
        "ตรวจประวัติรับชำระและประวัติการติดตามก่อนติดต่อหน่วยงานทุกครั้ง",
    ])

    document.add_page_break()
    add_section_heading(
        document,
        16,
        "ตรวจสอบและปิดสัญญา",
        "ปิดสัญญาเมื่อยอดเงินต้นและยอดตามตารางชำระครบถ้วนและตรวจสอบหลักฐานแล้ว",
    )
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    widths = (Inches(1.65), Inches(2.85), Inches(1.75))
    headers = ("ช่วงงาน", "สิ่งที่ต้องตรวจ", "ผู้รับผิดชอบหลัก")
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.text = header
        set_cell_shading(cell, GREEN)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_paragraph(paragraph, size=10.5, bold=True, color=WHITE, space_after=0, line_spacing=1.0)

    rows = (
        ("ก่อนครบกำหนด", "ตรวจงวดถัดไปและแจ้งเตือนหน่วยงานล่วงหน้า", "เจ้าหน้าที่ส่งเสริม"),
        ("เมื่อรับชำระ", "ตรวจหลักฐาน เลขที่ใบรับ และแยกเงินต้น ดอกเบี้ย ค่าปรับ", "เจ้าหน้าที่จังหวัด"),
        ("เมื่อค้างชำระ", "บันทึกผลติดต่อและกำหนดวันติดตามครั้งถัดไป", "ส่งเสริมและจังหวัด"),
        ("เมื่อชำระครบ", "ตรวจยอดทุกงวด ประวัติรับชำระ และเอกสารปิดสัญญา", "จังหวัดหรือ Admin"),
    )
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = width
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            format_paragraph(paragraph, size=10.5, bold=False, color=DARK, space_after=0, line_spacing=1.05)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(5)
    lead = note.add_run("ก่อนปิดสัญญา ")
    set_run_font(lead, size=12, bold=True, color=GREEN)
    body = note.add_run("ตรวจว่าทุกงวดมีสถานะชำระแล้วหรือยกเว้น ยอดเงินต้นคงเหลือเป็นศูนย์ และไม่มีงานติดตามค้างอยู่")
    set_run_font(body, size=12, bold=False, color=DARK)

    add_steps(document, [
        "เปิดหน้าติดตามหนี้และตรวจยอดคงค้างของทุกงวด",
        "ตรวจประวัติรับชำระ เลขที่ใบรับ และผู้บันทึกรายการ",
        "ตรวจประวัติการติดตามและงานที่กำหนดไว้ในอนาคต",
        "เมื่อข้อมูลครบ ให้เปลี่ยนสถานะจากชำระครบเป็นปิดสัญญา",
        "หากต้องแก้ข้อมูลหลังอนุมัติ ให้ Admin ดำเนินการและระบุเหตุผลเพื่อเก็บประวัติการเปลี่ยนแปลง",
    ])


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11.5)

    title = document.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.color.rgb = RGBColor.from_string(DARK)
    title.font.size = Pt(34)
    title.font.bold = True
    title.paragraph_format.left_indent = Inches(0)
    title.paragraph_format.right_indent = Inches(0)
    title.paragraph_format.first_line_indent = Inches(0)
    remove_style_border(title)

    heading = document.styles["Heading 1"]
    heading.font.name = FONT
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading.font.color.rgb = RGBColor.from_string(DARK)
    heading.font.size = Pt(28)
    heading.font.bold = True
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.left_indent = Inches(0)
    heading.paragraph_format.right_indent = Inches(0)
    heading.paragraph_format.first_line_indent = Inches(0)
    remove_style_border(heading)


def main():
    prepare_crops()
    document = Document(str(SOURCE))
    configure_styles(document)
    update_existing_content(document)
    append_tracking_sections(document)
    normalize_page_breaks(document)
    document.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
